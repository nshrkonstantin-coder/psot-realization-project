"""
Обработка Excel-файла и рассылка персональных email работникам.
Поддерживает: parse (разбор xlsx), send (отправка одной строки), track (трекинг открытия), status (получение статуса).
"""
import json
import os
import smtplib
import base64
import io
import re
import uuid
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from typing import Dict, Any, List, Optional


def get_db():
    import psycopg2
    return psycopg2.connect(os.environ['DATABASE_URL'])


def validate_email(email: str) -> bool:
    pattern = r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$'
    return bool(re.match(pattern, email.strip()))


def parse_excel(file_bytes: bytes) -> Dict[str, Any]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return {'error': 'Файл пустой'}
        headers = [str(h).strip() if h is not None else '' for h in rows[0]]
        data_rows = []
        for row in rows[1:]:
            row_data = {}
            for i, cell in enumerate(row):
                key = headers[i] if i < len(headers) else f'col_{i}'
                row_data[key] = str(cell).strip() if cell is not None else ''
            data_rows.append(row_data)
        return {'headers': headers, 'rows': data_rows}
    except Exception as e:
        return {'error': f'Ошибка чтения файла: {str(e)}'}


def build_email_html(row: Dict[str, str], include_columns: List[str], sender_display: str, track_id: str) -> str:
    rows_html = ''
    for col in include_columns:
        value = row.get(col, '')
        rows_html += f'<tr><td style="padding:8px 12px;border:1px solid #ddd;font-weight:600;background:#f5f5f5;width:40%">{col}</td><td style="padding:8px 12px;border:1px solid #ddd;">{value}</td></tr>'

    track_url = f"https://functions.poehali.dev/2dab48c9-57c0-4f55-90e7-d93b326a6891?action=track&id={track_id}"
    pixel = f'<img src="{track_url}" width="1" height="1" style="display:none" alt="" />'

    return f"""<html><body style="font-family:Arial,sans-serif;color:#333;max-width:600px;margin:0 auto">
<div style="background:linear-gradient(135deg,#1e293b,#334155);padding:24px;border-radius:8px 8px 0 0">
  <h2 style="color:#fff;margin:0;font-size:20px">{sender_display}</h2>
  <p style="color:#94a3b8;margin:4px 0 0">Информационное сообщение</p>
</div>
<div style="padding:24px;border:1px solid #e2e8f0;border-top:none;border-radius:0 0 8px 8px">
  <p style="margin-top:0">Уважаемый сотрудник,<br>направляем вам следующую информацию:</p>
  <table style="width:100%;border-collapse:collapse;margin:16px 0">{rows_html}</table>
  <hr style="border:none;border-top:1px solid #e2e8f0;margin:20px 0">
  <p style="color:#64748b;font-size:13px;margin:0">{sender_display}</p>
</div>
{pixel}
</body></html>"""


def handler(event: Dict[str, Any], context: Any) -> Dict[str, Any]:
    """
    Excel-рассылка: parse (разбор файла), send (одна строка), track (пиксель открытия), status (статусы)
    """
    cors = {
        'Access-Control-Allow-Origin': '*',
        'Access-Control-Allow-Methods': 'GET, POST, OPTIONS',
        'Access-Control-Allow-Headers': 'Content-Type, X-User-Id, X-Auth-Token',
    }

    method = event.get('httpMethod', 'GET')

    if method == 'OPTIONS':
        return {'statusCode': 200, 'headers': cors, 'body': ''}

    # ── GET: трекинг-пиксель ──
    if method == 'GET':
        params = event.get('queryStringParameters') or {}
        track_id = params.get('id', '')
        if params.get('action') == 'track' and track_id:
            try:
                db = get_db()
                cur = db.cursor()
                cur.execute(
                    "UPDATE excel_mail_log SET status='opened', opened_at=NOW() WHERE track_id=%s AND status!='opened'",
                    (track_id,)
                )
                db.commit()
                cur.close()
                db.close()
            except Exception:
                pass
            gif_b64 = 'R0lGODlhAQABAIAAAAAAAP///yH5BAEAAAAALAAAAAABAAEAAAIBRAA7'
            return {
                'statusCode': 200,
                'headers': {**cors, 'Content-Type': 'image/gif', 'Cache-Control': 'no-cache'},
                'body': gif_b64,
                'isBase64Encoded': True
            }
        return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'error': 'Неизвестный запрос'})}

    if method != 'POST':
        return {'statusCode': 405, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'error': 'Только POST/GET'})}

    try:
        body = json.loads(event.get('body', '{}'))
    except Exception:
        return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'error': 'Неверный JSON'})}

    action = body.get('action', 'parse')

    # ── parse ──
    if action == 'parse':
        file_b64 = body.get('file_base64', '')
        if not file_b64:
            return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': 'Файл не передан'})}
        try:
            file_bytes = base64.b64decode(file_b64)
        except Exception:
            return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': 'Ошибка декодирования файла'})}
        result = parse_excel(file_bytes)
        if 'error' in result:
            return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': result['error']})}
        return {'statusCode': 200, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'success': True, 'headers': result['headers'], 'rows': result['rows']})}

    # ── send (одна строка) ──
    elif action == 'send':
        headers_list: List[str] = body.get('headers', [])
        row: Dict[str, str] = body.get('row', {})
        sender_display: str = body.get('sender_name', 'АСУБТ').strip() or 'АСУБТ'
        subject: str = body.get('subject', 'Информационное сообщение').strip() or 'Информационное сообщение'
        sender_user_id: Optional[int] = body.get('user_id')

        email_col: Optional[str] = None
        include_col: Optional[str] = None
        for h in headers_list:
            hl = h.lower().strip()
            if 'электронная почта' in hl or hl in ('email', 'e-mail', 'почта'):
                email_col = h
            if 'включить в рассылку' in hl or 'рассылка' in hl:
                include_col = h

        if not email_col:
            return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': 'Не найдена колонка «Электронная почта»'})}

        to_email = row.get(email_col, '').strip()
        if not to_email or not validate_email(to_email):
            return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': f'Неверный email: {to_email}'})}

        exclude = {include_col, email_col} if include_col else {email_col}
        include_columns = [h for h in headers_list if h not in exclude]

        track_id = uuid.uuid4().hex

        smtp_host = os.environ.get('SMTP_HOST')
        smtp_port = int(os.environ.get('SMTP_PORT', '587'))
        smtp_user = os.environ.get('SMTP_USER')
        smtp_password = (os.environ.get('SMTP_PASSWORD_NEW') or
                         os.environ.get('YANDEX_SMTP_PASSWORD') or
                         os.environ.get('SMTP_PASSWORD'))

        if not all([smtp_host, smtp_user, smtp_password]):
            return {'statusCode': 500, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': 'SMTP не настроен'})}

        html = build_email_html(row, include_columns, sender_display, track_id)
        msg = MIMEMultipart('alternative')
        msg['Subject'] = subject
        msg['From'] = f'{sender_display} <{smtp_user}>'
        msg['To'] = to_email
        msg.attach(MIMEText(html, 'html', 'utf-8'))

        try:
            smtp = smtplib.SMTP(smtp_host, smtp_port, timeout=30)
            smtp.ehlo()
            if smtp_port == 587:
                smtp.starttls()
                smtp.ehlo()
            smtp.login(smtp_user, smtp_password)
            smtp.send_message(msg)
            smtp.quit()
        except smtplib.SMTPAuthenticationError:
            return {'statusCode': 500, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': 'Ошибка аутентификации SMTP'})}
        except Exception as e:
            return {'statusCode': 500, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': str(e)})}

        try:
            db = get_db()
            cur = db.cursor()
            cur.execute(
                "INSERT INTO excel_mail_log (track_id, sender_user_id, recipient_email, subject, sender_display, status, row_data) VALUES (%s, %s, %s, %s, %s, 'sent', %s)",
                (track_id, sender_user_id, to_email, subject, sender_display, json.dumps(row))
            )
            db.commit()
            cur.close()
            db.close()
        except Exception:
            pass

        return {'statusCode': 200, 'headers': {**cors, 'Content-Type': 'application/json'},
                'body': json.dumps({'success': True, 'track_id': track_id, 'email': to_email})}

    # ── status ──
    elif action == 'status':
        track_ids: List[str] = body.get('track_ids', [])
        if not track_ids:
            return {'statusCode': 200, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'success': True, 'statuses': {}})}
        try:
            db = get_db()
            cur = db.cursor()
            placeholders = ','.join(['%s'] * len(track_ids))
            cur.execute(
                f"SELECT track_id, status, sent_at, opened_at FROM excel_mail_log WHERE track_id IN ({placeholders})",
                track_ids
            )
            rows_db = cur.fetchall()
            cur.close()
            db.close()
            statuses = {}
            for r in rows_db:
                statuses[r[0]] = {
                    'status': r[1],
                    'sent_at': r[2].isoformat() if r[2] else None,
                    'opened_at': r[3].isoformat() if r[3] else None,
                }
            return {'statusCode': 200, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'success': True, 'statuses': statuses})}
        except Exception as e:
            return {'statusCode': 500, 'headers': {**cors, 'Content-Type': 'application/json'},
                    'body': json.dumps({'error': str(e)})}

    return {'statusCode': 400, 'headers': {**cors, 'Content-Type': 'application/json'},
            'body': json.dumps({'error': 'Неизвестный action'})}